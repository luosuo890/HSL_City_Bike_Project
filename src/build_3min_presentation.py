"""
Generator script for HSL_Bike_Heatmap_Presentation_3min.pptx.
Creates a professional, 6-slide, 3-minute course presentation deck
aligned with HSL_City_Bike_Project_Canvas.md, reports/findings.md, and verified figures.
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml

# Color Palette Constants
COLOR_DARK_BG = "12313A"      # Slate Teal (Slide 1 background)
COLOR_LIGHT_BG = "F7FAFA"     # Off-White / Soft Ice (Slides 2-6 background)
COLOR_CARD_BG = "FFFFFF"      # Pure white card background
COLOR_CARD_BORDER = "DCE6E8"  # Subtle grey-teal border
COLOR_PRIMARY_DARK = RGBColor(0x12, 0x31, 0x3A)  # Main text / titles on light
COLOR_PRIMARY_LIGHT = RGBColor(0xFF, 0xFF, 0xFF) # Main text on dark
COLOR_TEAL_ACCENT = RGBColor(0x00, 0x7E, 0x87)   # Primary brand accent
COLOR_MUTED_SLATE = RGBColor(0x54, 0x70, 0x78)   # Secondary / subtext
COLOR_LIGHT_ACCENT = RGBColor(0xA6, 0xDC, 0xDD)  # Accent on dark
COLOR_LIGHT_MUTED = RGBColor(0xC8, 0xDB, 0xDD)   # Secondary text on dark
COLOR_CORAL_ACCENT = RGBColor(0xD4, 0x5C, 0x39)  # Highlight / deficit color
COLOR_BLUE_ACCENT = RGBColor(0x2B, 0x5C, 0x8F)   # Secondary blue accent

FONT_HEADING = "Aptos"
FONT_BODY = "Aptos"


def set_slide_background(slide, color_hex):
    """Sets solid background color via DrawingML XML."""
    bg_xml = f"""
    <p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
          xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
      <p:bgPr>
        <a:solidFill>
          <a:srgbClr val="{color_hex}"/>
        </a:solidFill>
        <a:effectLst/>
      </p:bgPr>
    </p:bg>
    """
    cSld = slide._element.cSld
    existing_bgs = cSld.xpath("p:bg")
    for b in existing_bgs:
        cSld.remove(b)
    cSld.insert(0, parse_xml(bg_xml))


def add_slide_header(slide, slide_num, category, title, subtitle, is_dark=False):
    """Adds a standardized top header banner and bottom-right page indicator."""
    # Category badge / topic tracker
    tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10.5), Inches(0.35))
    tf_cat = tb_cat.text_frame
    tf_cat.word_wrap = True
    tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = f"{slide_num:02d} | {category.upper()}"
    p_cat.font.name = FONT_HEADING
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_LIGHT_ACCENT if is_dark else COLOR_TEAL_ACCENT

    # Main Slide Title
    tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.0), Inches(0.55))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_PRIMARY_LIGHT if is_dark else COLOR_PRIMARY_DARK

    # Subtitle / Core Takeaway
    tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.0), Inches(0.4))
    tf_sub = tb_sub.text_frame
    tf_sub.word_wrap = True
    tf_sub.margin_left = tf_sub.margin_right = tf_sub.margin_top = tf_sub.margin_bottom = 0
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = subtitle
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(13)
    p_sub.font.color.rgb = COLOR_LIGHT_MUTED if is_dark else COLOR_MUTED_SLATE

    # Page indicator (bottom right)
    tb_num = slide.shapes.add_textbox(Inches(11.5), Inches(6.9), Inches(1.0), Inches(0.3))
    tf_num = tb_num.text_frame
    tf_num.margin_left = tf_num.margin_right = tf_num.margin_top = tf_num.margin_bottom = 0
    p_num = tf_num.paragraphs[0]
    p_num.alignment = PP_ALIGN.RIGHT
    p_num.text = f"{slide_num:02d} / 06"
    p_num.font.name = FONT_BODY
    p_num.font.size = Pt(11)
    p_num.font.bold = True
    p_num.font.color.rgb = RGBColor(0x9C, 0xBA, 0xBC) if is_dark else COLOR_MUTED_SLATE


def add_card(slide, left, top, width, height, bg_hex="FFFFFF", border_hex="DCE6E8"):
    """Adds a styled rounded card container for content grouping."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    r = int(bg_hex[0:2], 16)
    g = int(bg_hex[2:4], 16)
    b = int(bg_hex[4:6], 16)
    shape.fill.fore_color.rgb = RGBColor(r, g, b)
    if border_hex:
        shape.line.color.rgb = RGBColor(int(border_hex[0:2], 16), int(border_hex[2:4], 16), int(border_hex[4:6], 16))
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def set_speaker_notes(slide, timing, duration, notes_text):
    """Sets structured speaker notes on slide with auto-computed word count."""
    wc = len(notes_text.strip().split())
    formatted_notes = f"{timing} | {duration} | {wc} words\n\n{notes_text.strip()}"
    slide.notes_slide.notes_text_frame.text = formatted_notes


def build_presentation(base_pptx_path, out_pptx_path):
    prs = Presentation(base_pptx_path)

    # Remove existing 9 slides cleanly
    for sldId in list(prs.slides._sldIdLst):
        rId = sldId.rId
        prs.slides._sldIdLst.remove(sldId)
        prs.part.drop_rel(rId)

    blank_layout = prs.slide_layouts[5]

    # =========================================================================
    # SLIDE 1: Title & Research Problem (Dark Theme)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, COLOR_DARK_BG)
    add_slide_header(s1, 1, "Project Overview", "Seasonal & Spatial Dynamics of HSL City Bikes",
                     "Exploratory Demand Heatmaps and Rebalancing Benchmarks for Helsinki & Espoo", is_dark=True)

    # Question Banner
    tb_q = s1.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.73), Inches(1.3))
    tf_q = tb_q.text_frame
    tf_q.word_wrap = True
    tf_q.margin_left = tf_q.margin_right = tf_q.margin_top = tf_q.margin_bottom = 0
    p_q1 = tf_q.paragraphs[0]
    p_q1.text = "How does city-bike demand vary across stations, seasons, and hours—"
    p_q1.font.name = FONT_HEADING
    p_q1.font.size = Pt(26)
    p_q1.font.bold = True
    p_q1.font.color.rgb = COLOR_PRIMARY_LIGHT

    p_q2 = tf_q.add_paragraph()
    p_q2.text = "and how can operators triage rebalancing bottlenecks?"
    p_q2.font.name = FONT_HEADING
    p_q2.font.size = Pt(26)
    p_q2.font.bold = True
    p_q2.font.color.rgb = COLOR_LIGHT_ACCENT
    p_q2.space_before = Pt(4)

    # 3 Pillar Summary Cards (Dark Theme)
    pillars = [
        ("Macro Seasonality", "Two Complete Cycling Seasons (2024–2025)",
         "Evaluates 14 operating months from April 1 to October 31, mapping unimodal summer demand peaks and shoulder contraction."),
        ("Spatial Concentration", "Station Heatmaps & Commuter Tidal Swings",
         "Identifies core downtown rail hubs, radial commuter sinks, and dramatic morning vs. evening tidal directional reversals."),
        ("Predictive Benchmarks", "Out-of-Time Forecasting & Station Triage",
         "Compares historical same-month baseline vs. ML models under causal within-season lags to benchmark departure and return errors.")
    ]
    card_w = Inches(3.75)
    card_h = Inches(2.2)
    card_y = Inches(3.6)
    for i, (title, sub, desc) in enumerate(pillars):
        card_x = Inches(0.8 + i * 4.0)
        add_card(s1, card_x, card_y, card_w, card_h, bg_hex="19414D", border_hex="265A68")

        tb = s1.shapes.add_textbox(card_x + Inches(0.2), card_y + Inches(0.18), card_w - Inches(0.4), card_h - Inches(0.36))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_LIGHT_ACCENT

        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.name = FONT_BODY
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_PRIMARY_LIGHT
        p2.space_before = Pt(4)

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.name = FONT_BODY
        p3.font.size = Pt(10.5)
        p3.font.color.rgb = COLOR_LIGHT_MUTED
        p3.space_before = Pt(6)

    # Footer Metadata
    tb_foot = s1.shapes.add_textbox(Inches(0.8), Inches(6.1), Inches(10.5), Inches(0.5))
    tf_foot = tb_foot.text_frame
    tf_foot.word_wrap = True
    p_auth = tf_foot.paragraphs[0]
    p_auth.text = "Authors: Chenyu Wu  ·  Jizhou Wang  ·  Jussi Honkio"
    p_auth.font.name = FONT_BODY
    p_auth.font.size = Pt(13)
    p_auth.font.bold = True
    p_auth.font.color.rgb = COLOR_PRIMARY_LIGHT

    p_meta = tf_foot.add_paragraph()
    p_meta.text = "Canvas Workshop 3 Presentation  |  Target Speaking Time: 3 Minutes (180s)  |  Helsinki Region Transport (HSL)"
    p_meta.font.name = FONT_BODY
    p_meta.font.size = Pt(10.5)
    p_meta.font.color.rgb = COLOR_LIGHT_MUTED
    p_meta.space_before = Pt(2)

    set_speaker_notes(s1, "0:00–0:20", "20 seconds",
                      "Good morning. Our project investigates seasonal and spatial bike demand across Helsinki and Espoo to help operators triage rebalancing bottlenecks. Using two complete cycling seasons, we examine system-wide seasonality, station-level concentration heatmaps, commuter tidal reversals, and an out-of-time forecasting benchmark for departures and returns.")

    # =========================================================================
    # SLIDE 2: Multi-Season Empirical Pipeline (Light Theme)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, COLOR_LIGHT_BG)
    add_slide_header(s2, 2, "Empirical Dataset & Pipeline", "Multi-Season Empirical Pipeline (2024–2025)",
                     "4.95M Clean Journeys Across 457 Stations Integrated with FMI Weather Ground Truth")

    # 4 Key Stat Cards across the top
    stats = [
        ("4.95M", "Valid Trips", "97.04% retained from 5.10M raw trips across 2024 (2.51M) & 2025 (2.44M)"),
        ("457", "Official Stations", "Official ArcGIS layer with dock capacity & WGS84 GPS; 4 unmapped docks preserved"),
        ("14", "Season Months", "Two full cycling seasons (Apr 1 – Oct 31); Helsinki 88.8%, Espoo 11.2%"),
        ("100%", "Weather Coverage", "FMI daily temp & rain (428 days per station); strictly historical observations")
    ]
    stat_w = Inches(2.78)
    stat_h = Inches(1.8)
    stat_y = Inches(1.9)
    for i, (val, lbl, detail) in enumerate(stats):
        stat_x = Inches(0.8 + i * 2.98)
        add_card(s2, stat_x, stat_y, stat_w, stat_h, bg_hex="FFFFFF", border_hex="DCE6E8")

        tb = s2.shapes.add_textbox(stat_x + Inches(0.18), stat_y + Inches(0.15), stat_w - Inches(0.36), stat_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p_val = tf.paragraphs[0]
        p_val.text = val
        p_val.font.name = FONT_HEADING
        p_val.font.size = Pt(28)
        p_val.font.bold = True
        p_val.font.color.rgb = COLOR_TEAL_ACCENT

        p_lbl = tf.add_paragraph()
        p_lbl.text = lbl
        p_lbl.font.name = FONT_HEADING
        p_lbl.font.size = Pt(13)
        p_lbl.font.bold = True
        p_lbl.font.color.rgb = COLOR_PRIMARY_DARK
        p_lbl.space_before = Pt(2)

        p_det = tf.add_paragraph()
        p_det.text = detail
        p_det.font.name = FONT_BODY
        p_det.font.size = Pt(9.5)
        p_det.font.color.rgb = COLOR_MUTED_SLATE
        p_det.space_before = Pt(4)

    # Bottom Split: Left = Quality Filters, Right = Data Artifacts
    col_w = Inches(5.72)
    col_h = Inches(2.75)
    col_y = Inches(3.95)

    # Left Card: Data Quality Safeguards
    add_card(s2, Inches(0.8), col_y, col_w, col_h, bg_hex="FFFFFF", border_hex="DCE6E8")
    tb_l = s2.shapes.add_textbox(Inches(1.05), col_y + Inches(0.2), col_w - Inches(0.5), col_h - Inches(0.4))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "Rigorous Data Quality Safeguards"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_DARK

    filters = [
        ("Realistic Physical Trip Bounds: ", "Filtered trips by duration (10s ≤ t ≤ 24h), distance (≥ 10m), and speed (≤ 10 m/s; 36 km/h) to remove dock test-swaps and GPS tracking errors."),
        ("Europe/Helsinki Local Time: ", "Parsed all journeys in local Finnish time to preserve true daylight and commuter diurnal behavior across midnight crossings without UTC timezone shift."),
        ("Zero Silent Imputation: ", "Stations absent from official ArcGIS layers (4 in 2024, 7 in 2025) retained with explicit NaN coordinates to avoid spatial bias.")
    ]
    for b_title, b_desc in filters:
        p = tf_l.add_paragraph()
        p.space_before = Pt(6)
        r1 = p.add_run()
        r1.text = "• " + b_title
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = COLOR_PRIMARY_DARK
        r2 = p.add_run()
        r2.text = b_desc
        r2.font.size = Pt(10)
        r2.font.color.rgb = COLOR_MUTED_SLATE

    # Right Card: Analytical Artifacts
    add_card(s2, Inches(6.8), col_y, col_w, col_h, bg_hex="FFFFFF", border_hex="DCE6E8")
    tb_r = s2.shapes.add_textbox(Inches(7.05), col_y + Inches(0.2), col_w - Inches(0.5), col_h - Inches(0.4))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p = tf_r.paragraphs[0]
    p.text = "Analytical Tables & Testing Suite"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_DARK

    artifacts = [
        ("station_month.csv (6,395 rows): ", "Primary analytical unit: departures, returns, net flow, capacity, coordinates, and monthly FMI weather across 14 operating months."),
        ("station_intraday_monthly.csv: ", "Aggregated 270,647 station-month-hour-daytype records capturing weekday vs. weekend and rush-hour tidal swings."),
        ("Automated Pytest Suite (5/5 PASSED): ", "Validates FMI -1.0 mm trace rain convention, weather completeness, primary key uniqueness, and station mapping integrity.")
    ]
    for b_title, b_desc in artifacts:
        p = tf_r.add_paragraph()
        p.space_before = Pt(6)
        r1 = p.add_run()
        r1.text = "• " + b_title
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = COLOR_PRIMARY_DARK
        r2 = p.add_run()
        r2.text = b_desc
        r2.font.size = Pt(10)
        r2.font.color.rgb = COLOR_MUTED_SLATE

    set_speaker_notes(s2, "0:20–0:50", "30 seconds",
                      "Our analysis builds on 4.95 million cleaned trips across 2024 and 2025, retaining 97% of raw records after quality filtering. We integrated 457 official stations with GPS coordinates and capacity, alongside daily weather observations from FMI. All timestamps use Europe/Helsinki local time, unmapped stations are preserved without imputation, and observed weather was strictly excluded from out-of-time forecasting to prevent leakage. This establishes a robust, reproducible dataset for both municipal networks.")

    # =========================================================================
    # SLIDE 3: Station-Month Demand Intensity Heatmap (Light Theme)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, COLOR_LIGHT_BG)
    add_slide_header(s3, 3, "Spatial & Seasonal Demand Intensity", "Station-Month Demand Intensity Heatmap",
                     "Core Transit Anchors Dominate a Pronounced Unimodal Nordic Cycling Season")

    # Left: Embed Real Figure 2 (Station-Month Demand Heatmap)
    fig2_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "reports", "figures", "fig2_station_month_demand_heatmap.png")
    img_left = Inches(0.8)
    img_top = Inches(1.85)
    img_width = Inches(7.4)
    img_height = Inches(4.75)
    if os.path.exists(fig2_path):
        s3.shapes.add_picture(fig2_path, img_left, img_top, img_width, img_height)
    else:
        # Fallback card if image missing
        add_card(s3, img_left, img_top, img_width, img_height)

    # Right: 3 Structured Analytical Takeaways
    panel_left = Inches(8.45)
    panel_width = Inches(4.08)

    cards_s3 = [
        ("Unimodal Season Envelope",
         "Ridership follows subarctic climate, peaking in July/August (~450k–492k trips/mo) and contracting sharply in shoulder months (April: 162k–243k; October: 198k–207k)."),
        ("Heavy Transit Hub Concentration",
         "The Top 30 stations account for a disproportionate share of volume. Core anchors—Pasila railway station, Central Railway, Kalasatama, Itämerentori—sustain 5,000–9,000 departures/month throughout summer."),
        ("Descriptive Weather Co-Movement",
         "Daily volume correlates positively with mean temperature (r = +0.74) and negatively with precipitation (r = -0.25). Reflects seasonal daylight and holidays; strictly descriptive, non-causal.")
    ]
    card_h_s3 = Inches(1.45)
    for j, (head, body) in enumerate(cards_s3):
        cy = Inches(1.85 + j * 1.65)
        add_card(s3, panel_left, cy, panel_width, card_h_s3, bg_hex="FFFFFF", border_hex="DCE6E8")

        tb = s3.shapes.add_textbox(panel_left + Inches(0.18), cy + Inches(0.12), panel_width - Inches(0.36), card_h_s3 - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = head
        p.font.name = FONT_HEADING
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEAL_ACCENT

        p2 = tf.add_paragraph()
        p2.text = body
        p2.font.name = FONT_BODY
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = COLOR_PRIMARY_DARK
        p2.space_before = Pt(3)

    set_speaker_notes(s3, "0:50–1:30", "40 seconds",
                      "This station-month heatmap shows the Top 30 busiest stations across all 14 months. Demand follows a pronounced unimodal summer envelope, peaking in July and August at nearly 500,000 monthly trips before dropping sharply in autumn. Activity is heavily concentrated: Helsinki accounts for 89% of trips, anchored by major transit hubs like Pasila and Central Railway Station. Daily volume correlates positively with temperature and negatively with rain, reflecting seasonal co-movement rather than a direct causal mechanism.")

    # =========================================================================
    # SLIDE 4: Intraday Peak Asymmetry & Tidal Reversals (Light Theme)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, COLOR_LIGHT_BG)
    add_slide_header(s4, 4, "Intraday Travel Dynamics", "Commuter Peak Asymmetry & Tidal Flow Reversals",
                     "Weekday Evening Surge (+70% to +105%) and Polarized Morning-Evening Station Imbalances")

    col4_w = Inches(5.72)
    col4_h = Inches(4.75)
    col4_y = Inches(1.85)

    # Left Card: Evening Peak Surge & Diurnal Curves
    add_card(s4, Inches(0.8), col4_y, col4_w, col4_h, bg_hex="FFFFFF", border_hex="DCE6E8")
    tb_asym = s4.shapes.add_textbox(Inches(1.05), col4_y + Inches(0.2), col4_w - Inches(0.5), col4_h - Inches(0.4))
    tf_asym = tb_asym.text_frame
    tf_asym.word_wrap = True

    p = tf_asym.paragraphs[0]
    p.text = "The Evening Commute Surge (+70% to +105%)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_DARK

    p = tf_asym.add_paragraph()
    p.text = "Standardized 2-hour weekday windows (Morning: 07:00–09:00 vs. Evening: 16:00–18:00):"
    p.font.name = FONT_BODY
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_MUTED_SLATE
    p.space_before = Pt(3)

    metrics_asym = [
        ("Helsinki 2025: ", "Morning 1,207.5  →  Evening 2,154.4 trips/day (+78.4% surge in evening)"),
        ("Espoo 2025: ", "Morning 140.7  →  Evening 275.6 trips/day (+95.9% surge in evening)"),
        ("Behavioral Driver: ", "Morning trips are utilitarian home-to-work journeys. Evening trips combine homeward commutes with dining, retail, and social activities under extended northern daylight."),
        ("Diurnal Shape: ", "Weekdays show sharp bimodal peaks (08:00 & 17:00). Weekends follow a unimodal midday plateau (12:00–17:00) driven by recreational travel without morning rush.")
    ]
    for m_head, m_body in metrics_asym:
        p = tf_asym.add_paragraph()
        p.space_before = Pt(6)
        r1 = p.add_run()
        r1.text = "• " + m_head
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = COLOR_TEAL_ACCENT if "2025" in m_head else COLOR_PRIMARY_DARK
        r2 = p.add_run()
        r2.text = m_body
        r2.font.size = Pt(10)
        r2.font.color.rgb = COLOR_PRIMARY_DARK

    p_hypo = tf_asym.add_paragraph()
    p_hypo.space_before = Pt(8)
    p_hypo.text = "Operational Working Hypothesis: Pre-staging bicycle redistribution prior to the afternoon surge (e.g. before 15:30) could mitigate evening dock stockouts; verifying requires live dock telemetry."
    p_hypo.font.name = FONT_BODY
    p_hypo.font.size = Pt(9.5)
    p_hypo.font.italic = True
    p_hypo.font.color.rgb = COLOR_MUTED_SLATE

    # Right Card: Station Tidal Reversal Table
    add_card(s4, Inches(6.8), col4_y, col4_w, col4_h, bg_hex="FFFFFF", border_hex="DCE6E8")
    tb_tbl_title = s4.shapes.add_textbox(Inches(7.05), col4_y + Inches(0.2), col4_w - Inches(0.5), Inches(0.5))
    tf_tt = tb_tbl_title.text_frame
    tf_tt.word_wrap = True
    p = tf_tt.paragraphs[0]
    p.text = "Diurnal Commuter Tidal Reversal Matrix"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_DARK

    # Create Table on Right Card
    tbl_shape = s4.shapes.add_table(4, 4, Inches(7.05), col4_y + Inches(0.75), col4_w - Inches(0.5), Inches(2.2))
    tbl = tbl_shape.table
    tbl.columns[0].width = Inches(1.7)
    tbl.columns[1].width = Inches(1.15)
    tbl.columns[2].width = Inches(1.15)
    tbl.columns[3].width = Inches(1.22)

    headers = ["Station Name", "AM Peak (07-09)", "PM Peak (16-18)", "Commuter Role"]
    data = [
        ["Sompasaari", "-7,973 trips", "+2,808 trips", "Residential Feeder (AM sink, PM fill)"],
        ["Salmisaarenranta", "+7,343 trips", "-5,244 trips", "Workplace Hub (AM fill, PM sink)"],
        ["Pasila railway station", "-2,618 trips", "-2,607 trips", "Intermodal Rail Hub (All-day sink)"]
    ]
    for c_idx, h in enumerate(headers):
        cell = tbl.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x12, 0x31, 0x3A)
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.name = FONT_HEADING
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.alignment = PP_ALIGN.CENTER if c_idx > 0 else PP_ALIGN.LEFT

    for r_idx, row_data in enumerate(data):
        for c_idx, val in enumerate(row_data):
            cell = tbl.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xF0, 0xF6, 0xF7) if r_idx % 2 == 1 else RGBColor(0xFF, 0xFF, 0xFF)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.name = FONT_BODY
            p.font.size = Pt(9.5)
            if c_idx == 0:
                p.font.bold = True
                p.font.color.rgb = COLOR_PRIMARY_DARK
            elif "-" in val:
                p.font.bold = True
                p.font.color.rgb = COLOR_CORAL_ACCENT
                p.alignment = PP_ALIGN.RIGHT
            elif "+" in val:
                p.font.bold = True
                p.font.color.rgb = COLOR_TEAL_ACCENT
                p.alignment = PP_ALIGN.RIGHT
            else:
                p.font.color.rgb = COLOR_MUTED_SLATE
                p.alignment = PP_ALIGN.LEFT

    tb_note = s4.shapes.add_textbox(Inches(7.05), col4_y + Inches(3.1), col4_w - Inches(0.5), Inches(1.4))
    tf_note = tb_note.text_frame
    tf_note.word_wrap = True
    p = tf_note.paragraphs[0]
    p.text = "* Net Flow = Returns - Departures across all 2024 & 2025 weekday peak hours."
    p.font.name = FONT_BODY
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_DARK

    p2 = tf_note.add_paragraph()
    p2.text = "Critical Operational Distinction: Values represent cumulative passenger trip events, NOT current dock inventory or required truck workloads. Diurnal reversals naturally offset daytime flows; physical redistribution requires real-time 15-minute dock monitoring."
    p2.font.name = FONT_BODY
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = COLOR_MUTED_SLATE
    p2.space_before = Pt(3)

    set_speaker_notes(s4, "1:30–2:05", "35 seconds",
                      "Beyond monthly totals, intraday dynamics reveal strong commute asymmetry. Weekday evening peaks generate 70 to 105% more departures than morning peaks, as homeward trips blend with dining and leisure. Simultaneously, stations undergo tidal reversals: residential hubs like Sompasaari face heavy morning deficits but evening surpluses, while workplace hubs like Salmisaarenranta exhibit the opposite. These flows represent cumulative trip events, not real-time dock inventory or truck loads. Understanding this diurnal swing is essential for timing pre-staging operations.")

    # =========================================================================
    # SLIDE 5: Predictive Benchmarks & Station Rebalancing Triage (Light Theme)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, COLOR_LIGHT_BG)
    add_slide_header(s5, 5, "Forecasting Benchmarks & Triage", "Predictive Benchmarks & Station Rebalancing Triage",
                     "Historical Baseline Minimizes Volume Error; Net Flow Guides Rebalancing Screening")

    col5_w = Inches(5.72)
    col5_h = Inches(4.75)
    col5_y = Inches(1.85)

    # Left: Out-of-Time Model Benchmark (2025 Test Cohort)
    add_card(s5, Inches(0.8), col5_y, col5_w, col5_h, bg_hex="FFFFFF", border_hex="DCE6E8")
    tb_m = s5.shapes.add_textbox(Inches(1.05), col5_y + Inches(0.2), col5_w - Inches(0.5), col5_h - Inches(0.4))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True

    p = tf_m.paragraphs[0]
    p.text = "Model Comparison: 2024 Train → 2025 Test"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_DARK

    p = tf_m.add_paragraph()
    p.text = "Evaluated on strictly matched out-of-time test cohort (May–Oct 2025, N = 2,718):"
    p.font.name = FONT_BODY
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_MUTED_SLATE
    p.space_before = Pt(3)

    models_info = [
        ("Same-Month Baseline: ", "17.58% WAPE", " (MAE: 141.9 trips/mo)",
         "Decisively lowest volume error! Naturally aligns with solar irradiance, seasonal temperature, and academic calendar without transitional lag delay."),
        ("Decision Tree: ", "24.72% WAPE", " (MAE: 199.5 trips/mo)",
         "Segmented tree rules show moderate accuracy but exhibit high variance on station transitions."),
        ("Ridge Regression: ", "28.00% WAPE", " (MAE: 225.9 trips/mo)",
         "Severe relative errors on low-volume stations (Q1 WAPE 124.1%) due to rigid global month shifts; yet achieves 65%–80% Top-20 deficit overlap (ρ = 0.76).")
    ]
    for m_name, wape, mae, m_desc in models_info:
        p = tf_m.add_paragraph()
        p.space_before = Pt(6)
        r1 = p.add_run()
        r1.text = "• " + m_name
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = COLOR_PRIMARY_DARK

        r2 = p.add_run()
        r2.text = wape
        r2.font.bold = True
        r2.font.size = Pt(11)
        r2.font.color.rgb = COLOR_TEAL_ACCENT if "17.58" in wape else COLOR_PRIMARY_DARK

        r3 = p.add_run()
        r3.text = mae
        r3.font.size = Pt(10)
        r3.font.color.rgb = COLOR_MUTED_SLATE

        p_desc = tf_m.add_paragraph()
        p_desc.text = "   " + m_desc
        p_desc.font.size = Pt(9.5)
        p_desc.font.color.rgb = COLOR_MUTED_SLATE
        p_desc.space_before = Pt(1)

    p_m_note = tf_m.add_paragraph()
    p_m_note.space_before = Pt(8)
    p_m_note.text = "* Unit: MAE is measured in completed trips per station-month, NOT physical bicycles. Single out-of-time test cycle (2025); exploratory benchmark without claims of statistical significance."
    p_m_note.font.name = FONT_BODY
    p_m_note.font.size = Pt(9)
    p_m_note.font.italic = True
    p_m_note.font.color.rgb = COLOR_MUTED_SLATE

    # Right: Station Rebalancing Priority Triage
    add_card(s5, Inches(6.8), col5_y, col5_w, col5_h, bg_hex="FFFFFF", border_hex="DCE6E8")
    tb_triage = s5.shapes.add_textbox(Inches(7.05), col5_y + Inches(0.2), col5_w - Inches(0.5), col5_h - Inches(0.4))
    tf_tr = tb_triage.text_frame
    tf_tr.word_wrap = True

    p = tf_tr.paragraphs[0]
    p.text = "Station Rebalancing Priority Triage (Net Flow)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY_DARK

    p = tf_tr.add_paragraph()
    p.text = "Cumulative net flow (returns - departures) across all 14 months highlights systemic bottlenecks:"
    p.font.name = FONT_BODY
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_MUTED_SLATE
    p.space_before = Pt(3)

    deficit_stations = [
        ("Pasila railway station (ID 113): ", "-11,014 trips", " (Cap: 40)", "Major regional train interchange sink; passengers disembark from commuter trains and cycle away to surrounding offices."),
        ("Mäkelänkatu (ID 120): ", "-5,275 trips", " (Cap: 28)", "Arterial transit corridor sink; chronic morning & afternoon outbound trips."),
        ("Opastinsilta (ID 203): ", "-5,049 trips", " (Cap: 40)", "East Pasila commercial/administrative complex; rapid bike exhaustion."),
        ("Karhupuisto (ID 43): ", "-4,421 trips", " (Cap: 24)", "Kallio residential hub; heavy outbound departures downhill toward center.")
    ]
    for s_name, net_f, cap, s_desc in deficit_stations:
        p = tf_tr.add_paragraph()
        p.space_before = Pt(5)
        r1 = p.add_run()
        r1.text = "• " + s_name
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_PRIMARY_DARK

        r2 = p.add_run()
        r2.text = net_f
        r2.font.bold = True
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = COLOR_CORAL_ACCENT

        r3 = p.add_run()
        r3.text = cap
        r3.font.size = Pt(9.5)
        r3.font.color.rgb = COLOR_MUTED_SLATE

        p_desc = tf_tr.add_paragraph()
        p_desc.text = "   " + s_desc
        p_desc.font.size = Pt(9)
        p_desc.font.color.rgb = COLOR_MUTED_SLATE
        p_desc.space_before = Pt(1)

    p_tr_note = tf_tr.add_paragraph()
    p_tr_note.space_before = Pt(8)
    p_tr_note.text = "* Operational Guardrail: Rankings provide an empirical triage list for planning attention, NOT rigid van dispatch instructions. Rebalancing requires telemetry on dock fullness."
    p_tr_note.font.name = FONT_BODY
    p_tr_note.font.size = Pt(9)
    p_tr_note.font.italic = True
    p_tr_note.font.color.rgb = COLOR_MUTED_SLATE

    set_speaker_notes(s5, "2:05–2:40", "35 seconds",
                      "Evaluating on 2025 data, a historical same-month baseline achieved the lowest volume error at 17.6% WAPE, outperforming Decision Trees and Ridge regression because it naturally captures Nordic seasonal shifts. However, Ridge provided strong rank ordering for early-season deficit screening. Concurrently, cumulative net flow identifies persistent transit sinks like Pasila railway station. These rankings provide an empirical triage list for planning, not rigid dispatch commands. No single model dominates all operational objectives.")

    # =========================================================================
    # SLIDE 6: Methodological Guardrails & Operational Conclusions (Light Theme)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, COLOR_LIGHT_BG)
    add_slide_header(s6, 6, "Methodological Guardrails & Conclusions", "Methodological Guardrails & Operational Conclusions",
                     "Responsible Interpretation of Preliminary Empirical Benchmarks to Guide Fleet Operations")

    # 4 Guardrails 2x2 Grid across Top
    guardrails = [
        ("Weather = Retrospective Observation",
         "Weather parameters reflect historical FMI observations, strictly excluded from 2025 predictive models. Descriptive correlations (r ≈ 0.74 with temp) co-move with subarctic daylight (19h in June vs. 9h in Oct) and holiday calendars; they do NOT establish direct causality."),
        ("Forecast Error Unit: Trips ≠ Bikes",
         "Mean Absolute Error (MAE = 141.9 trips/mo) measures completed trip events per station-month, NOT a physical bicycle shortage. In bike-sharing networks, individual bikes circulate across multiple trips; trip error does not equal missing fleet stock."),
        ("Macro Net Flow ≠ Rebalancing Workloads",
         "Cumulative net flow (returns - departures) reflects 30-day directional passenger drift. Diurnal tidal reversals partially offset daytime imbalances. Executing physical truck dispatch requires real-time 15-minute dock fullness telemetry, not monthly sums."),
        ("Single Out-of-Time Test Season (2025)",
         "Covering two operating seasons yields a single out-of-time test cycle. These empirical benchmarks represent exploratory baselines, NOT multi-year confidence intervals. No claims of statistical significance or proven long-term superiority are made.")
    ]
    gw = Inches(5.72)
    gh = Inches(1.85)
    for k, (g_title, g_desc) in enumerate(guardrails):
        gx = Inches(0.8) if k % 2 == 0 else Inches(6.8)
        gy = Inches(1.85) if k < 2 else Inches(3.85)
        add_card(s6, gx, gy, gw, gh, bg_hex="FFFFFF", border_hex="DCE6E8")

        tb = s6.shapes.add_textbox(gx + Inches(0.2), gy + Inches(0.12), gw - Inches(0.4), gh - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = g_title
        p.font.name = FONT_HEADING
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_CORAL_ACCENT if "≠" in g_title else COLOR_TEAL_ACCENT

        p2 = tf.add_paragraph()
        p2.text = g_desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = COLOR_PRIMARY_DARK
        p2.space_before = Pt(3)

    # Bottom Full-Width Synthesis Card
    bot_y = Inches(5.82)
    bot_w = Inches(11.72)
    bot_h = Inches(0.95)
    add_card(s6, Inches(0.8), bot_y, bot_w, bot_h, bg_hex="12313A", border_hex="12313A")

    tb_syn = s6.shapes.add_textbox(Inches(1.0), bot_y + Inches(0.1), bot_w - Inches(0.4), bot_h - Inches(0.2))
    tf_syn = tb_syn.text_frame
    tf_syn.word_wrap = True

    p = tf_syn.paragraphs[0]
    p.text = "Operational Synthesis for HSL Planners:"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_LIGHT_ACCENT

    p2 = tf_syn.add_paragraph()
    p2.text = "1) Use Historical Same-Month Baseline for macro seasonal procurement and fleet allocation.  2) Reference Net Flow triage to prioritize structural transit sinks (Pasila corridor).  3) Connect real-time dock occupancy before executing operational van rebalancing."
    p2.font.name = FONT_BODY
    p2.font.size = Pt(10)
    p2.font.color.rgb = COLOR_PRIMARY_LIGHT
    p2.space_before = Pt(2)

    set_speaker_notes(s6, "2:40–3:00", "20 seconds",
                      "To conclude, our findings establish clear operational guardrails. Weather is retrospective observation, forecast errors are in trip units rather than physical bicycles, and net flows reflect passenger drift. With a single 2025 test season, these preliminary benchmarks provide an empirical foundation before connecting real-time dock telemetry. Thank you, and we welcome your questions.")

    prs.save(out_pptx_path)
    print(f"Successfully generated: {out_pptx_path}")
    print(f"Total slides: {len(prs.slides)}")


if __name__ == "__main__":
    project_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    base_file = os.path.join(project_dir, "HSL_Bike_Heatmap_Presentation.pptx")
    out_file = os.path.join(project_dir, "HSL_Bike_Heatmap_Presentation_3min.pptx")
    build_presentation(base_file, out_file)
